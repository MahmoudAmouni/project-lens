"""Read a Word (.docx) file into pages of text plus tables, and decide its status.

Word files have no fixed pages, so pages follow the page breaks Word saved in the file:
the breaks from its last layout when there are any, else the hard page breaks."""

import logging
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table as WordTable

from Server.ingestion.models import STATUS_CORRUPT, STATUS_EMPTY, STATUS_ENCRYPTED, Document, Page, Table
from Server.ingestion.money_parser import currency_in_notes
from Server.ingestion.pdf_reader import clean_text
from Server.ingestion.table_cleaner import clean_table

logger = logging.getLogger("projectlens")

OLE_SIGNATURE = b"\xd0\xcf\x11\xe0"  # password protected .docx files are saved in this old binary format
PARAGRAPH, TABLE = qn("w:p"), qn("w:tbl")
TEXT, TAB, LINE_BREAK, RENDERED_BREAK = qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:lastRenderedPageBreak")
BREAK_TYPE = qn("w:type")


class PageCounter:
    """Collects text page by page while walking the document in order."""

    def __init__(self, use_rendered_breaks: bool):
        self.use_rendered_breaks = use_rendered_breaks
        self.number = 1
        self.texts: dict[int, list[str]] = {1: []}

    def is_page_break(self, element) -> bool:
        if self.use_rendered_breaks:
            return element.tag == RENDERED_BREAK
        return element.tag == LINE_BREAK and element.get(BREAK_TYPE) == "page"

    def next_page(self) -> None:
        self.number += 1
        self.texts.setdefault(self.number, [])

    def add_paragraph(self, paragraph) -> None:
        """Add one paragraph's text; a break inside it moves the rest to the next page."""
        line: list[str] = []
        for element in paragraph.iter(TEXT, TAB, LINE_BREAK, RENDERED_BREAK):
            if self.is_page_break(element):
                self.add_line("".join(line))
                line = []
                self.next_page()
            elif element.tag == TEXT:
                line.append(element.text or "")
            elif element.tag == TAB:
                line.append("\t")
            elif element.tag == LINE_BREAK:
                line.append("\n")
        self.add_line("".join(line))

    def add_line(self, text: str) -> None:
        if text.strip():
            self.texts[self.number].append(text)

    def skip_breaks(self, element) -> None:
        """Count the page breaks inside a table row (its text is added separately)."""
        for child in element.iter(LINE_BREAK, RENDERED_BREAK):
            if self.is_page_break(child):
                self.next_page()

    def pages(self) -> list[Page]:
        return [Page(number=number, text=clean_text("\n".join(lines))) for number, lines in self.texts.items()]


def read_word_table(table: WordTable, counter: PageCounter, number: int) -> Table | None:
    """Rows of the table as text, plus the page each row is on. Its text also goes into the page text."""
    note_currency = currency_in_notes(counter.texts[counter.number])  # text above the table on its page
    rows, row_pages = [], []
    for row in table.rows:
        cells = [clean_text(cell.text) or None for cell in row.cells]  # merged cells repeat their value
        rows.append(cells)
        row_pages.append(counter.number)
        counter.add_line("\t".join(cell or "" for cell in cells))
        counter.skip_breaks(row._tr)
    cleaned = clean_table(rows, note_currency=note_currency)
    if cleaned is None:
        return None
    data = cleaned.data
    page_column = "Page" if "Page" not in data.columns else "Source page"
    data.insert(0, page_column, [row_pages[row_number - 1] for row_number in data.index])
    first, last = row_pages[0], row_pages[-1]
    return Table(
        id=f"page_{first}_table_{number}",
        data=data,
        location=f"page {first}" if first == last else f"pages {first}-{last}",
        currencies=cleaned.currencies,
        summary_rows=cleaned.summary_rows,
    )


def read_word(path: Path, document: Document) -> Document:
    """Fill `document` with pages, tables and status. Never raises."""
    with path.open("rb") as file:
        if file.read(4) == OLE_SIGNATURE:
            document.status = STATUS_ENCRYPTED
            document.message = "This Word file is password protected or in the old .doc format."
            return document
    try:
        word = docx.Document(str(path))
    except Exception as error:
        document.status = STATUS_CORRUPT
        document.message = "This Word file is damaged or is not a real Word file."
        logger.warning("Could not open %s: %s", path.name, error)
        return document

    body = word.element.body
    counter = PageCounter(use_rendered_breaks=next(body.iter(RENDERED_BREAK), None) is not None)
    tables_per_page: dict[int, int] = {}
    for element in body.iterchildren():
        if element.tag == PARAGRAPH:
            counter.add_paragraph(element)
        elif element.tag == TABLE:
            tables_per_page[counter.number] = tables_per_page.get(counter.number, 0) + 1
            table = read_word_table(WordTable(element, word), counter, tables_per_page[counter.number])
            if table is not None:
                document.tables.append(table)

    document.pages = counter.pages()
    if not any(page.text for page in document.pages) and not document.tables:
        document.status = STATUS_EMPTY
        document.message = "This Word file has no text."
    return document
