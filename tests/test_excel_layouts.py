"""Real-world Excel layouts: title rows merged across the columns above the header must not become the header."""

from pathlib import Path

import openpyxl

from Server.ingestion.excel_reader import read_excel
from Server.ingestion.models import Document

HEADER = ["Ref", "Description of Work", "Unit", "Qty", "Rate (SAR)", "Amount (SAR)"]


def write_boq_with_merged_titles(path: Path) -> None:
    """Like many tender BOQs: three merged title rows, a blank row, the header, sections and subtotals."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Civil BOQ"
    for title in ["RIYADH METRO DEPOT - PACKAGE 2", "BILL OF QUANTITIES - CIVIL WORKS", "Revision C - 12 June 2025"]:
        sheet.append([title])
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=6)
    sheet.append([])
    sheet.append(HEADER)
    sheet.append([None, "SECTION A - EARTHWORKS"])
    sheet.append(["A1", "Bulk excavation in rock", "m3", 31500, 42, 1323000])
    sheet.append(["A2", "Engineered backfill", "m3", 22750, 27.5, "625,625.00"])
    sheet.append([None, "Sub-Total Section A", None, None, None, 1948625])
    sheet.append([None, "GRAND TOTAL", None, None, None, 1948625])
    workbook.save(path)


def read(path: Path) -> Document:
    return read_excel(path, Document(name=path.name, path=str(path), file_type="excel"))


def test_merged_title_rows_are_not_taken_as_the_header(tmp_path):
    path = tmp_path / "BOQ_Civil_Works.xlsx"
    write_boq_with_merged_titles(path)
    table = read(path).tables[0]
    assert list(table.data.columns) == HEADER
    assert table.currencies == {"Rate (SAR)": "SAR", "Amount (SAR)": "SAR"}
    assert table.data.loc[7, "Amount (SAR)"] == 1323000  # row numbers stay the Excel row numbers
    assert table.data.loc[8, "Amount (SAR)"] == 625625  # an amount typed as text is still a number
    assert table.summary_rows == [9, 10]  # sub-total and grand total are left out of sums


def test_a_real_header_with_repeated_words_is_still_found(tmp_path):
    path = tmp_path / "Repeated.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Amount", "Amount", "Description"])
    workbook.active.append([100, 200, "Steel"])
    workbook.save(path)
    assert list(read(path).tables[0].data.columns) == ["Amount", "Amount (2)", "Description"]
