"""Every sample file gets the right status, and tables come out clean and exact."""

from pathlib import Path

import pytest

from Server.ingestion.folder_loader import load_folder
from Server.ingestion.folder_scanner import FolderNotFoundError, scan_folder
from Server.ingestion.money_parser import detect_currency, parse_number
from Server.ingestion.table_cleaner import clean_table
from tests.sample_content import (
    BOQ_PHASE1_ITEMS,
    BOQ_PHASE2_AMOUNTS,
    ELECTRICAL_AMOUNTS,
)

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_docs"


@pytest.fixture(scope="module")
def documents():
    result = load_folder(SAMPLE_DIR)
    return {document.name: document for document in result.documents}


@pytest.mark.parametrize(
    ("name", "status"),
    [
        ("Contract_A.pdf", "ok"),
        ("BOQ_Phase1.xlsx", "ok"),
        ("BOQ_Phase2.xlsx", "ok"),
        ("BOQ_Electrical.pdf", "ok"),
        ("Technical_Specifications.pdf", "ok"),
        ("Arabic_Scope.pdf", "ok"),
        ("Scanned_Drawing.pdf", "scanned"),
        ("Corrupt_Report.pdf", "corrupt"),
        ("Protected_Annex.pdf", "encrypted"),
        ("Empty_Workbook.xlsx", "empty"),
        ("Zero_Bytes.xlsx", "empty"),
        ("Meeting_Notes.docx", "ok"),
    ],
)
def test_each_sample_file_gets_the_right_status(documents, name, status):
    assert documents[name].status == status
    if status != "ok":
        assert documents[name].message


def test_unsupported_files_are_skipped_not_errors():
    _, skipped = scan_folder(SAMPLE_DIR)
    assert [file.name for file in skipped] == ["Site_Photo.png"]


def test_missing_folder_raises_clear_error():
    with pytest.raises(FolderNotFoundError, match="does not exist"):
        scan_folder(SAMPLE_DIR / "no_such_folder")


def test_large_file_is_marked_too_large():
    result = load_folder(SAMPLE_DIR, max_file_mb=0.01)
    statuses = {document.name: document.status for document in result.documents}
    assert statuses["Technical_Specifications.pdf"] == "too_large"


def test_pdf_pages_are_one_based_and_complete(documents):
    spec = documents["Technical_Specifications.pdf"]
    assert len(spec.pages) == 220
    assert spec.pages[136].number == 137
    assert "curing period of 7 days" in spec.pages[136].text


def test_arabic_pdf_text_uses_normal_letters(documents):
    text = documents["Arabic_Scope.pdf"].pages[0].text
    assert "نطاق العمل" in text


def test_excel_boq_keeps_exact_amounts_and_row_numbers(documents):
    table = documents["BOQ_Phase1.xlsx"].find_table("boq")
    amounts = table.data["Total Amount (SAR)"]
    assert amounts.loc[2] == BOQ_PHASE1_ITEMS[0][5]  # data starts on Excel row 2
    assert table.summary_rows == [len(BOQ_PHASE1_ITEMS) + 2]
    items = amounts.drop(index=table.summary_rows)
    assert items.sum() == sum(item[5] for item in BOQ_PHASE1_ITEMS)
    assert table.currencies["Total Amount (SAR)"] == "SAR"


def test_header_on_row_four_and_merged_cells_are_handled(documents):
    table = documents["BOQ_Phase2.xlsx"].find_table("BOQ Phase 2")
    assert list(table.data.columns) == [
        "Item No.", "Section", "Description", "Unit", "Quantity", "Amount", "Cell note",
    ]  # fmt: skip
    assert table.data["Cell note"].dropna().tolist() == ["Amount: Lump sum"]  # text kept, not lost as a blank
    assert table.data.index[0] == 5
    assert table.data.loc[7, "Section"] == "Mechanical"  # filled from merged cell B5:B7
    assert table.data.loc[5, "Item No."] == "4.1"  # item codes stay text


def test_money_text_becomes_numbers(documents):
    table = documents["BOQ_Phase2.xlsx"].find_table("BOQ Phase 2")
    item_amounts = table.data["Amount"].drop(index=table.summary_rows).tolist()
    for got, expected in zip(item_amounts, BOQ_PHASE2_AMOUNTS, strict=True):
        assert (got != got) if expected is None else got == expected  # NaN for "Lump sum"


def test_usd_sheet_is_marked_as_usd(documents):
    table = documents["BOQ_Phase2.xlsx"].find_table("Imported Equipment")
    assert table.currencies == {"Price (USD)": "USD"}


def test_pdf_table_over_two_pages_is_one_table_with_pages(documents):
    tables = documents["BOQ_Electrical.pdf"].tables
    assert len(tables) == 1
    table = tables[0]
    assert table.location == "pages 2-3"
    assert table.data["Amount (SAR)"].tolist() == ELECTRICAL_AMOUNTS
    assert set(table.data["Page"]) == {2, 3}


@pytest.mark.parametrize(
    ("text", "number"),
    [
        ("SAR 1,200,000", 1200000),
        ("SAR 1,250,000.00", 1250000),
        ("1.25M", 1250000),
        ("2.5 million", 2500000),
        ("(500)", -500),
        ("٣٬٥٠٠٫٥", 3500.5),
        ("1 250 000", 1250000),
        ("1.250.000,50", 1250000.5),
        ("ر.س 2,000", 2000),
        ("150 m", None),  # meters, not millions
        ("20 mm", None),  # millimetres (aggregate size), not millions
        ("20mm", None),
        ("2026-01-05", None),
        ("Lump sum", None),
        ("-", None),
    ],
)
def test_parse_number(text, number):
    assert parse_number(text) == number


def test_detect_currency():
    assert detect_currency("Amount (SAR)") == "SAR"
    assert detect_currency("Price (USD)") == "USD"
    assert detect_currency("SARS outbreak") is None


def test_clean_table_returns_none_for_blank_grid():
    assert clean_table([[None, ""], ["  ", None]]) is None


def test_big_excel_fast_path_keeps_row_numbers(tmp_path, monkeypatch):
    import openpyxl

    import Server.ingestion.excel_reader as excel_reader

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["B4"], sheet["C4"] = "Item", "Amount (SAR)"
    sheet["B5"], sheet["C5"] = "1.1", "SAR 1,200,000"
    sheet["B7"], sheet["C7"] = "1.2", 300
    workbook.save(tmp_path / "big.xlsx")
    monkeypatch.setattr(excel_reader, "FAST_MODE_MB", -1)  # force the calamine path
    table = load_folder(tmp_path).documents[0].tables[0]
    assert table.data["Amount (SAR)"].to_dict() == {5: 1200000, 7: 300}


def sheet_rows(*blocks):
    """Rows of a sheet: each block is a list of rows; None stands for one blank row."""
    return [[None, None, None] if row is None else row for block in blocks for row in block]


def test_two_tables_in_one_sheet_are_split_with_exact_rows():
    from Server.ingestion.excel_reader import sheet_to_tables

    rows = sheet_rows(
        [["Item", "Description", "Amount (SAR)"], ["1", "Steel", 500], ["2", "Concrete", 700]],
        [None, None],
        [["Ref", "Equipment", "Price (USD)"], ["E-1", "Lift", 90]],
    )
    tables = sheet_to_tables("BOQ", rows)
    assert [t.id for t in tables] == ["BOQ (table 1)", "BOQ (table 2)"]
    assert tables[1].data.index.tolist() == [7]  # Excel row 7
    assert tables[1].currencies == {"Price (USD)": "USD"}


def test_gap_between_sections_of_one_table_does_not_split_it():
    from Server.ingestion.excel_reader import sheet_to_tables

    rows = sheet_rows(
        [["Bill of Quantities", None, None]],
        [None, None],
        [["Item", "Description", "Amount (SAR)"], ["1", "Steel", 500]],
        [None, None],
        [["2", "Concrete", 700]],
    )
    tables = sheet_to_tables("BOQ", rows)
    assert [t.id for t in tables] == ["BOQ"]
    assert tables[0].data.index.tolist() == [5, 8]
