"""Title lines above a sheet's header ("Revision A (SUPERSEDED)") reach the model: two versions can be told apart."""

from pathlib import Path

import openpyxl
import pytest

from AI.tools.describe_table import describe_table
from AI.tools.list_documents import document_overview, list_documents
from AI.tools.tool_context import ToolContext
from Server.ingestion.folder_loader import load_folder

HEADER = ["Milestone", "Description", "Amount (SAR)"]


def write_schedule(path: Path, title: str, amount: int) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Milestones"
    sheet.append([title])
    sheet.append(HEADER)
    sheet.append([1, "Advance payment", amount])
    sheet.append([2, "Foundations complete", amount * 2])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


@pytest.fixture()
def context(tmp_path):
    write_schedule(tmp_path / "Archive" / "Payment_Schedule.xlsx", "Payment schedule - Revision A (SUPERSEDED)", 4600)
    write_schedule(tmp_path / "Tender" / "Payment_Schedule.xlsx", "Payment schedule - Revision C", 4875)
    return ToolContext(documents=load_folder(tmp_path).documents)


def test_the_title_line_is_kept_on_the_table(context):
    notes = {document.name: document.tables[0].notes for document in context.documents}
    assert notes["Archive/Payment_Schedule.xlsx"] == ["Payment schedule - Revision A (SUPERSEDED)"]
    assert notes["Tender/Payment_Schedule.xlsx"] == ["Payment schedule - Revision C"]


def test_the_header_and_rows_are_not_title_lines(context):
    table = context.documents[0].tables[0]
    assert list(table.data.columns) == HEADER and len(table.data) == 2


def test_the_model_sees_the_titles_in_every_tool(context):
    overview = document_overview(context.documents)
    assert "[Payment schedule - Revision A (SUPERSEDED)]" in overview
    assert "[Payment schedule - Revision C]" in overview
    described = describe_table(context, "Archive/Payment_Schedule.xlsx", "Milestones").data
    assert described["title_lines"] == ["Payment schedule - Revision A (SUPERSEDED)"]
    listed = list_documents(context).data["documents"]
    assert any(t["title_lines"] == ["Payment schedule - Revision C"] for d in listed for t in d["tables"])


def test_a_sheet_without_a_title_has_no_title_lines(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.append(HEADER)
    workbook.active.append([1, "Advance payment", 100])
    workbook.save(tmp_path / "Plain.xlsx")
    document = load_folder(tmp_path).documents[0]
    assert document.tables[0].notes == []
    assert "[" not in document_overview([document])
