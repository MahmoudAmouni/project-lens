"""Create the test files in sample_docs/, including the broken ones.

Run: python -m tests.make_sample_docs
"""

from pathlib import Path

import docx
import openpyxl
import pymupdf

from tests.sample_content import (
    ARABIC_SCOPE_HTML,
    BOQ_PHASE1_ITEMS,
    BOQ_PHASE2_ITEMS,
    CONTRACT_PAGES,
    ELECTRICAL_ITEMS,
    IMPORTED_EQUIPMENT_ITEMS,
    long_spec_page,
)

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_docs"
WINDOWS_FONTS = Path("C:/Windows/Fonts")


def write_text_pdf(path: Path, pages: list[str]) -> None:
    """A normal PDF, one string per page."""
    pdf = pymupdf.open()
    for text in pages:
        page = pdf.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 545, 800), text, fontsize=10)
    pdf.save(path)


def write_contract(path: Path) -> None:
    write_text_pdf(path, CONTRACT_PAGES)


def write_protected_contract(path: Path) -> None:
    pdf = pymupdf.open()
    pdf.new_page().insert_text((50, 80), "Confidential contract annex.")
    pdf.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="1234", owner_pw="owner-1234")


def write_long_spec(path: Path, page_count: int = 220) -> None:
    write_text_pdf(path, [long_spec_page(n) for n in range(1, page_count + 1)])


def write_scanned(path: Path) -> None:
    """Pages that are only pictures of text, like a scan."""
    source = pymupdf.open()
    source.new_page().insert_text((50, 80), "Site drawing A-101, scanned copy", fontsize=20)
    picture = source[0].get_pixmap(dpi=72)
    pdf = pymupdf.open()
    for _ in range(2):
        pdf.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842), pixmap=picture)
    pdf.save(path, deflate=True, garbage=3)


def draw_table_page(page: pymupdf.Page, header: list[str], rows: list[list[str]], widths: list[int]) -> None:
    """Draw a ruled table (cell borders) so pdfplumber can find it."""
    top, height = 60, 22
    for row_index, row in enumerate([header, *rows]):
        left = 40
        for text, width in zip(row, widths, strict=True):
            cell = pymupdf.Rect(left, top + row_index * height, left + width, top + (row_index + 1) * height)
            page.draw_rect(cell, color=(0, 0, 0), width=0.5)
            page.insert_textbox(cell + (3, 5, -3, 0), text, fontsize=8)
            left += width


def write_electrical_boq(path: Path) -> None:
    """A BOQ table in a PDF that runs over two pages (header repeated)."""
    header = ["Item", "Description", "Unit", "Qty", "Amount (SAR)"]
    widths = [45, 230, 45, 60, 130]
    pdf = pymupdf.open()
    intro = pdf.new_page()
    intro.insert_text((50, 80), "Electrical works - Bill of Quantities", fontsize=16)
    intro.insert_text((50, 110), "All amounts in Saudi Riyals (SAR), excluding VAT.", fontsize=10)
    half = len(ELECTRICAL_ITEMS) // 2
    for part in (ELECTRICAL_ITEMS[:half], ELECTRICAL_ITEMS[half:]):
        draw_table_page(pdf.new_page(), header, part, widths)
    pdf.save(path)


def write_arabic_scope(path: Path) -> bool:
    """An Arabic PDF. Needs a font with Arabic letters (Arial on Windows)."""
    if not (WINDOWS_FONTS / "arial.ttf").exists():
        return False
    pdf = pymupdf.open()
    page = pdf.new_page()
    css = "@font-face {font-family: ar; src: url(arial.ttf);} * {font-family: ar; font-size: 13px;}"
    page.insert_htmlbox(
        page.rect + (50, 50, -50, -50),
        ARABIC_SCOPE_HTML,
        css=css,
        archive=pymupdf.Archive(str(WINDOWS_FONTS)),
    )
    pdf.subset_fonts()
    pdf.save(path, deflate=True, garbage=3)
    return True


def write_boq_phase1(path: Path) -> None:
    """Clean BOQ: header on row 1, real numbers, a total row at the end."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "BOQ"
    sheet.append(["Item", "Description", "Unit", "Qty", "Unit Rate (SAR)", "Total Amount (SAR)"])
    for item in BOQ_PHASE1_ITEMS:
        sheet.append(list(item))
    sheet.append([None, "Grand Total", None, None, None, sum(i[5] for i in BOQ_PHASE1_ITEMS)])
    summary = workbook.create_sheet("Summary")
    summary.append(["Section", "Amount (SAR)"])
    summary.append(["Phase 1 civil and structural works", sum(i[5] for i in BOQ_PHASE1_ITEMS)])
    workbook.save(path)


def write_boq_phase2(path: Path) -> None:
    """Messy BOQ: title rows, header on row 4, merged cells, money as text, a USD sheet."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "BOQ Phase 2"
    sheet["A1"] = "Bill of Quantities - Phase 2 (MEP and Finishes)"
    sheet["A2"] = "Project: Riyadh Office Tower"
    sheet.append([])
    sheet.append(["Item No.", "Section", "Description", "Unit", "Quantity", "Amount"])
    for item in BOQ_PHASE2_ITEMS:
        sheet.append(list(item))
    sheet.merge_cells("B5:B7")  # "Mechanical" section spans three rows
    sheet.merge_cells("B8:B10")
    sheet.append(["", "Sub Total", None, None, None, "SAR 9,885,000"])
    usd = workbook.create_sheet("Imported Equipment")
    usd.append(["Ref", "Equipment", "Qty", "Price (USD)"])
    for item in IMPORTED_EQUIPMENT_ITEMS:
        usd.append(list(item))
    workbook.save(path)


def write_empty_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.create_sheet("Sheet2")
    workbook.save(path)


def write_meeting_notes(path: Path) -> None:
    """A real Word file: two pages (hard page break) and an action table on page 2."""
    notes = docx.Document()
    notes.add_heading("Site Meeting No. 7", level=1)
    notes.add_paragraph("Attendees: Client representative, Main Contractor, MEP Subcontractor.")
    notes.add_paragraph("The contractor reported that slab casting for Level 3 is complete.")
    notes.add_page_break()
    notes.add_heading("Action items", level=2)
    table = notes.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = "Action", "Owner", "Status"
    for action, owner, status in [
        ("Submit revised shop drawings for ducts", "MEP Subcontractor", "Open"),
        ("Approve facade mock-up", "Client representative", "Closed"),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = action, owner, status
    notes.save(path)


def main() -> None:
    SAMPLE_DIR.mkdir(exist_ok=True)
    write_contract(SAMPLE_DIR / "Contract_A.pdf")
    write_protected_contract(SAMPLE_DIR / "Protected_Annex.pdf")
    write_long_spec(SAMPLE_DIR / "Technical_Specifications.pdf")
    write_scanned(SAMPLE_DIR / "Scanned_Drawing.pdf")
    write_electrical_boq(SAMPLE_DIR / "BOQ_Electrical.pdf")
    write_boq_phase1(SAMPLE_DIR / "BOQ_Phase1.xlsx")
    write_boq_phase2(SAMPLE_DIR / "BOQ_Phase2.xlsx")
    write_empty_workbook(SAMPLE_DIR / "Empty_Workbook.xlsx")
    (SAMPLE_DIR / "Corrupt_Report.pdf").write_text("This is a text file renamed to .pdf\n")
    (SAMPLE_DIR / "Zero_Bytes.xlsx").write_bytes(b"")
    write_meeting_notes(SAMPLE_DIR / "Meeting_Notes.docx")
    (SAMPLE_DIR / "Site_Photo.png").write_bytes(b"\x89PNG\r\n\x1a\n")  # not a document: it should be skipped
    if not write_arabic_scope(SAMPLE_DIR / "Arabic_Scope.pdf"):
        print("Skipped Arabic_Scope.pdf: no Arabic font found.")
    print(f"Sample documents written to {SAMPLE_DIR}")


if __name__ == "__main__":
    main()
