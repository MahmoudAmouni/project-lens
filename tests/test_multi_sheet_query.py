"""query_table on a workbook with several sheets: one call covers every sheet, and "the largest" is found across all."""

import openpyxl
import pytest

from AI.tools.query_table import query_table
from AI.tools.tool_context import ToolContext, ToolError
from Server.ingestion.folder_loader import load_folder

BILLS = {
    "Preliminaries": [("PRE-1", "Site office", 2_500_000), ("PRE-2", "Fencing", 900_000)],
    "Civil": [("CIV-1", "Raft foundation", 8_000_000), ("CIV-2", "Columns", 3_000_000), ("CIV-3", "Kerbs", 400_000)],
    "Mechanical": [("MEC-1", "Chillers", 2_100_000), ("MEC-2", "Ducts", 1_200_000)],
}
OVER_2M = [{"column": "amount", "operator": ">", "value": 2_000_000}]


@pytest.fixture(scope="module")
def context(tmp_path_factory):
    folder = tmp_path_factory.mktemp("tender")
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Bill", "Amount (SAR)"])
    for name, rows in BILLS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(["Item", "Description", "Amount (SAR)"])
        for row in rows:
            sheet.append(list(row))
        summary.append([name, sum(amount for *_, amount in rows)])
    workbook.save(folder / "BOQ_Full_Package.xlsx")
    schedule = openpyxl.Workbook()  # another file with an amount column: it must stay out of a one-file query
    schedule.active.append(["Milestone", "Amount (SAR)"])
    schedule.active.append([1, 9_000_000])
    schedule.save(folder / "Payment_Schedule.xlsx")
    return ToolContext(documents=load_folder(folder).documents)


def test_a_file_without_a_sheet_searches_every_sheet_of_that_file_only(context):
    result = query_table(context, document="BOQ_Full_Package", filters=OVER_2M)
    assert result.data["total_match_count"] == 4  # PRE-1, CIV-1, CIV-2, MEC-1
    assert result.data["combined_sum_by_currency"] == {"SAR": 15_600_000}
    searched = {table["table_id"] for table in result.data["tables"]}
    assert searched == {"Preliminaries", "Civil", "Mechanical"}  # not Summary, not the payment schedule
    assert result.data["summary_sheets_left_out"] == ["BOQ_Full_Package.xlsx · Summary"]


def test_the_largest_item_is_ranked_across_all_sheets(context):
    result = query_table(context, document="BOQ_Full_Package.xlsx", sort_by="amount", limit=3)
    best = result.data["best_rows_of_all_tables"]
    assert [row["Item"] for row in best] == ["CIV-1", "CIV-2", "PRE-1"]
    assert best[0]["table_id"] == "Civil" and best[0]["Amount (SAR)"] == 8_000_000
    saved = context.saved_tables[result.data["table_result_id"]]
    assert len(saved["rows"]) == 3  # the overall top 3, not the top 3 of every sheet


def test_one_sheet_can_still_be_chosen(context):
    result = query_table(context, document="BOQ_Full_Package.xlsx", table_id="Mechanical", filters=OVER_2M)
    assert result.data["total_match_count"] == 1


def test_columns_only_on_the_summary_sheet_search_the_summary(context):
    result = query_table(
        context,
        document="BOQ_Full_Package.xlsx",
        filters=[{"column": "bill", "operator": "contains", "value": "civil"}],
    )
    assert [table["table_id"] for table in result.data["tables"]] == ["Summary"]


def test_a_wrong_column_in_one_file_names_its_sheets(context):
    with pytest.raises(ToolError, match="Tables: Summary, Preliminaries, Civil, Mechanical"):
        query_table(
            context, document="BOQ_Full_Package.xlsx", filters=[{"column": "weight", "operator": ">", "value": 1}]
        )
