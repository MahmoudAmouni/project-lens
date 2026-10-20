"""A currency named only in a note above the table ("All prices in US Dollars") reaches its money columns."""

from pathlib import Path

import openpyxl
import pymupdf
import pytest

from AI.tools.list_documents import document_overview
from Server.ingestion.excel_reader import read_excel
from Server.ingestion.models import Document
from Server.ingestion.money_parser import currency_in_notes, detect_currency
from Server.ingestion.pdf_reader import read_pdf
from tests.make_sample_docs import draw_table_page


@pytest.mark.parametrize(
    ("text", "code"),
    [("All prices in US Dollars", "USD"), ("Saudi Riyals", "SAR"), ("بالدولار الأمريكي", "USD"),
     ("Rate (SAR)", "SAR"), ("Supply of SARS masks", None), ("Description", None)],
)  # fmt: skip
def test_plural_and_arabic_currency_words_are_found(text, code):
    assert detect_currency(text) == code


def test_notes_must_name_exactly_one_currency():
    assert currency_in_notes(["Note: all prices are in US Dollars", "Revision C"]) == "USD"
    assert currency_in_notes(["Prices in USD", "Freight in SAR"]) is None
    assert currency_in_notes(["Revision C"]) is None


def write_sheet(path: Path, rows: list[list]) -> Document:
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    return read_excel(path, Document(name=path.name, path=str(path), file_type="excel"))


def test_a_note_above_the_header_gives_every_table_on_the_sheet_its_currency(tmp_path):
    document = write_sheet(tmp_path / "MEP.xlsx", [
        ["Note: all prices in this workbook are in US Dollars and exclude VAT."], [],
        ["Item No.", "Description", "Qty", "Unit Price", "Total Price"],
        ["M-01", "Chillers", 2, 285000, 570000], ["M-02", "Pumps", 1, 126000, 126000], [], [],
        ["Code", "Spare Part", "Qty", "Price Each", "Line Total"], ["SP-1", "Compressor", 1, 64000, 64000],
    ])  # fmt: skip
    main, spares = document.tables
    assert main.currencies == {"Unit Price": "USD", "Total Price": "USD"}  # not Qty
    assert spares.currencies == {"Price Each": "USD", "Line Total": "USD"}


def test_a_currency_in_the_header_wins_over_the_note(tmp_path):
    document = write_sheet(tmp_path / "Mixed.xlsx", [
        ["All prices in US Dollars"], ["Item", "Description", "Amount (SAR)"], ["A1", "Fence", 1000],
    ])  # fmt: skip
    assert document.tables[0].currencies == {"Amount (SAR)": "SAR"}


def test_without_a_note_the_overview_says_currency_not_stated(tmp_path):
    document = write_sheet(tmp_path / "Plain.xlsx", [["Item", "Total Price"], ["A1", 500], ["A2", 700]])
    document.status = "ok"
    assert document.tables[0].currencies == {}
    assert "currency not stated" in document_overview([document])


def test_a_note_on_the_pdf_page_gives_the_table_its_currency(tmp_path):
    path = tmp_path / "Quote.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((50, 40), "All prices are in US Dollars.", fontsize=10)
    draw_table_page(page, ["Item", "Description", "Total Price"], [["1", "Chiller", "570000"], ["2", "Pump", "126000"]],
                    [60, 250, 120])  # fmt: skip
    pdf.save(path)
    document = read_pdf(path, Document(name=path.name, path=str(path), file_type="pdf"))
    assert document.tables[0].currencies == {"Total Price": "USD"}
