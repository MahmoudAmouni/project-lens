"""Word files: pages follow Word's page breaks, tables can be cited by page, bad files get a status."""

from pathlib import Path

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement

from Server.ingestion.models import Document
from Server.ingestion.word_reader import read_word

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_docs"


def read(path: Path) -> Document:
    return read_word(path, Document(name=path.name, path=str(path), file_type="word"))


def test_sample_meeting_notes_has_two_pages_and_a_cited_table():
    document = read(SAMPLE_DIR / "Meeting_Notes.docx")
    assert document.status == "ok"
    assert [page.number for page in document.pages] == [1, 2]
    assert "slab casting for Level 3" in document.pages[0].text
    assert "Approve facade mock-up" in document.pages[1].text
    table = document.tables[0]
    assert table.id == "page_2_table_1" and table.location == "page 2"
    assert list(table.data.columns) == ["Page", "Action", "Owner", "Status"]
    assert table.data["Page"].tolist() == [2, 2]


def test_page_break_inside_a_paragraph_splits_its_text(tmp_path):
    path = tmp_path / "split.docx"
    word = docx.Document()
    paragraph = word.add_paragraph("before break")
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run("after break")
    word.save(path)
    pages = read(path).pages
    assert [(page.number, page.text) for page in pages] == [(1, "before break"), (2, "after break")]


def test_word_layout_breaks_win_over_hard_breaks(tmp_path):
    """A file saved by Word holds its real page ends; then hard breaks are not counted twice."""
    path = tmp_path / "rendered.docx"
    word = docx.Document()
    word.add_paragraph("page one")
    run = word.add_paragraph().add_run("page two")
    run._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))
    word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)  # ignored: not where Word ended the page
    word.save(path)
    assert [page.number for page in read(path).pages] == [1, 2]


def test_empty_word_file_is_marked_empty(tmp_path):
    path = tmp_path / "empty.docx"
    docx.Document().save(path)
    document = read(path)
    assert document.status == "empty" and document.message


def test_fake_word_file_is_marked_corrupt(tmp_path):
    path = tmp_path / "fake.docx"
    path.write_text("just text")
    document = read(path)
    assert document.status == "corrupt" and document.message


def test_password_protected_word_file_is_marked_encrypted(tmp_path):
    path = tmp_path / "locked.docx"
    path.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 100)
    document = read(path)
    assert document.status == "encrypted" and document.message
